import csv
import io
from datetime import date
from unittest import mock

import requests
from django.test import TestCase, override_settings

from core import services
from core.models import DeviceList, EventLog, NotificationLog, SensorLog

DEVICE_ID = "2069691213314072577"


def http(status=200, json_data=None):
    response = mock.Mock()
    response.status_code = status
    response.ok = 200 <= status < 300
    response.reason = "OK" if response.ok else "Error"
    response.json.return_value = json_data
    response.text = ""
    return response


def token_ok():
    return http(json_data={"code": "00000000", "data": {"access_token": "tok"}})


def event(event_id, event_type="in", target="Cross Line", event_time="2026-10-01 10:46:27"):
    return {
        "id": event_id,
        "eventTime": event_time,
        "eventLogVO": {"trackId": "1", "height": "157", "eventType": event_type, "recognitionTarget": target},
    }


def page(items):
    return http(json_data={"code": "00000000", "data": {"data": items, "page": 1, "pageSize": 10}})


@override_settings(ZK_CLIENT_ID="cid", ZK_CLIENT_SECRET="secret")
class SyncDeviceTests(TestCase):
    def setUp(self):
        self.device = DeviceList.objects.get(id=DEVICE_ID)
        self.device.baseline_done = True
        self.device.maximum_trigger = 3
        self.device.save()

    def run_sync(self, get_responses, post_responses=None):
        posts = [token_ok()] + list(post_responses or [])
        with mock.patch.object(services.requests, "get", side_effect=get_responses) as get, \
                mock.patch.object(services.requests, "post", side_effect=posts) as post:
            services.sync_device(self.device, services.ZKClient())
        self.device.refresh_from_db()
        return get, post

    def test_token_body_uses_env_credentials(self):
        _, post = self.run_sync([page([])])
        self.assertEqual(
            post.call_args_list[0].kwargs["json"],
            {"client_id": "cid", "client_secret": "secret", "grant_type": "client_credentials"},
        )

    def test_first_sync_stores_baseline_without_counting(self):
        self.device.baseline_done = False
        self.device.save()
        self.run_sync([page([event("a"), event("b")])])
        self.assertTrue(self.device.baseline_done)
        self.assertEqual(self.device.current_count, 0)
        self.assertEqual(EventLog.objects.count(), 2)

    def test_counts_only_in_cross_line(self):
        items = [
            event("1", "in"), event("2", "out"), event("3", "passby"),
            event("4", "passby and in"), event("5", "turnback"), event("6", "in", target="Body"),
        ]
        self.run_sync([page(items)])
        self.assertEqual(self.device.current_count, 1)
        self.assertEqual(list(EventLog.objects.filter(counted=True).values_list("id", flat=True).order_by("id")), ["1"])

    def test_duplicate_events_are_not_counted_twice(self):
        self.run_sync([page([event("1")])])
        self.run_sync([page([event("2"), event("1")])])
        self.assertEqual(self.device.current_count, 2)
        self.assertEqual(EventLog.objects.count(), 2)

    def test_pages_until_known_event(self):
        EventLog.objects.create(id="old", time="2026-10-01T10:00:00+07:00", device=self.device)
        full_page = [event(str(i)) for i in range(10)]
        self.device.maximum_trigger = 100
        self.device.save()
        get, _ = self.run_sync([page(full_page), page([event("x"), event("old")])])
        self.assertEqual(get.call_count, 2)
        self.assertEqual(self.device.current_count, 11)

    def test_threshold_dispatches_work_order_and_resets(self):
        _, post = self.run_sync(
            [page([event("1"), event("2"), event("3")])],
            [http(json_data={"status": "success", "wo_id": "WO-000001"})],
        )
        self.assertEqual(self.device.current_count, 0)
        log = NotificationLog.objects.get()
        self.assertEqual(log.response_status, "200 OK")
        self.assertEqual(post.call_args_list[1].args[0], "http://127.0.0.1:8000/dummy/api_iot.php")

    def test_inactive_real_endpoint_is_never_called(self):
        _, post = self.run_sync(
            [page([event("1"), event("2"), event("3")])],
            [http(json_data={"status": "success"})],
        )
        urls = [call.args[0] for call in post.call_args_list[1:]]
        self.assertEqual(urls, ["http://127.0.0.1:8000/dummy/api_iot.php"])

    def test_below_threshold_does_not_dispatch(self):
        self.run_sync([page([event("1"), event("2")])])
        self.assertEqual(self.device.current_count, 2)
        self.assertFalse(NotificationLog.objects.exists())

    def test_failed_work_order_is_logged_and_count_still_resets(self):
        self.run_sync(
            [page([event("1"), event("2"), event("3")])],
            [requests.ConnectionError("down")],
        )
        self.assertEqual(self.device.current_count, 0)
        self.assertEqual(NotificationLog.objects.get().response_status, "ERROR")

    def test_rejected_token_is_refreshed_and_retried(self):
        unauthorized = http(401, {"code": "A0230", "message": "token expired"})
        with mock.patch.object(services.requests, "get", side_effect=[unauthorized, page([event("1")])]), \
                mock.patch.object(services.requests, "post", side_effect=[token_ok(), token_ok()]) as post:
            services.sync_device(self.device, services.ZKClient())
        self.device.refresh_from_db()
        self.assertEqual(post.call_count, 2)
        self.assertEqual(self.device.current_count, 1)
        self.assertEqual(SensorLog.objects.filter(status=SensorLog.STATUS_OFFLINE).count(), 1)

    def test_same_day_events_keep_counting(self):
        self.device.current_count = 1
        self.device.count_date = date(2026, 10, 1)
        self.device.save()
        self.run_sync([page([event("1")])])
        self.assertEqual(self.device.current_count, 2)
        self.assertEqual(self.device.count_date, date(2026, 10, 1))

    def test_new_day_resets_count(self):
        self.device.current_count = 2
        self.device.count_date = date(2026, 9, 30)
        self.device.save()
        self.run_sync([page([event("1")])])
        self.assertEqual(self.device.current_count, 1)
        self.assertEqual(self.device.count_date, date(2026, 10, 1))

    def test_batch_across_midnight_discards_previous_day_without_work_order(self):
        self.device.current_count = 2
        self.device.count_date = date(2026, 9, 30)
        self.device.save()
        # Newest first, as ZK returns them: yesterday's 23:59 event reaches the threshold of 3.
        self.run_sync([page([
            event("today", event_time="2026-10-01 00:01:00"),
            event("yesterday", event_time="2026-09-30 23:59:00"),
        ])])
        self.assertEqual(self.device.current_count, 1)
        self.assertEqual(self.device.count_date, date(2026, 10, 1))
        self.assertFalse(NotificationLog.objects.exists())
        self.assertEqual(EventLog.objects.filter(counted=True).count(), 2)

    def test_empty_count_date_is_set_without_reset(self):
        self.device.current_count = 1
        self.device.save()
        self.run_sync([page([event("1")])])
        self.assertEqual(self.device.current_count, 2)
        self.assertEqual(self.device.count_date, date(2026, 10, 1))

    def test_ignored_event_on_new_day_does_not_reset(self):
        self.device.current_count = 2
        self.device.count_date = date(2026, 9, 30)
        self.device.save()
        self.run_sync([page([event("1", "out")])])
        self.assertEqual(self.device.current_count, 2)
        self.assertEqual(self.device.count_date, date(2026, 9, 30))

    def test_sensor_log_never_stores_access_token(self):
        self.run_sync([page([])])
        for log in SensorLog.objects.all():
            self.assertNotIn("tok", str(log.response))


class DailyRecapAdminTests(TestCase):
    def test_recap_shows_daily_totals(self):
        from django.contrib.auth.models import User

        device = DeviceList.objects.get(id=DEVICE_ID)
        EventLog.objects.create(id="a", time="2026-10-01T10:00:00+07:00", device=device, counted=True)
        EventLog.objects.create(id="b", time="2026-10-01T11:00:00+07:00", device=device, counted=False)
        NotificationLog.objects.create(device=device, endpoint_url="x", response_status="200 OK")
        NotificationLog.objects.create(device=device, endpoint_url="x", response_status="ERROR")
        self.client.force_login(User.objects.create_superuser("admin", "", "pw"))

        response = self.client.get("/admin/core/dailyrecap/")
        self.assertEqual(response.status_code, 200)
        rows = {row["day"].isoformat(): row for row in response.context["rows"]}
        self.assertEqual((rows["2026-10-01"]["total"], rows["2026-10-01"]["counted"]), (2, 1))
        today = max(rows)
        self.assertEqual((rows[today]["sent"], rows[today]["success"], rows[today]["failed"]), (2, 1, 1))


class ExportTests(TestCase):
    def setUp(self):
        from django.contrib.auth.models import User

        self.device = DeviceList.objects.get(id=DEVICE_ID)
        self.client.force_login(User.objects.create_superuser("admin", "", "pw"))

    def export_csv(self, url):
        """Submit the django-import-export export form (CSV, all fields) and return the CSV rows."""
        form = self.client.get(url).context["form"]
        fields = [name for name in form.fields if name not in ("format", "resource", "export_items")]
        data = {"format": "0", "resource": "0", **{name: "on" for name in fields}}
        response = self.client.post(url, data)
        self.assertEqual(response.status_code, 200)
        return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))

    def test_log_changelists_show_export_button(self):
        for model in ("eventlog", "sensorlog", "notificationlog"):
            response = self.client.get(f"/admin/core/{model}/")
            self.assertContains(response, f"/admin/core/{model}/export/")

    def test_event_log_export_respects_filters_and_uses_wib(self):
        EventLog.objects.create(id="a", time="2026-10-01T10:00:00+07:00", device=self.device,
                                event_type="in", recognition_target="Cross Line", counted=True)
        EventLog.objects.create(id="b", time="2026-10-01T11:00:00+07:00", device=self.device, event_type="out")
        rows = self.export_csv("/admin/core/eventlog/export/?event_type=in")
        self.assertEqual(rows[0], ["id", "time", "device", "event_type", "recognition_target",
                                   "track_id", "height", "counted"])
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][:4], ["a", "2026-10-01 10:00:00", DEVICE_ID, "in"])

    def test_notification_log_export_has_json_text(self):
        NotificationLog.objects.create(device=self.device, endpoint_url="x", response_status="200 OK",
                                       body={"REQ_DESC": "Toilet"}, head={}, response={"status": "success"})
        rows = self.export_csv("/admin/core/notificationlog/export/")
        record = dict(zip(rows[0], rows[1]))
        self.assertEqual(record["body"], '{"REQ_DESC": "Toilet"}')
        self.assertEqual(record["response"], '{"status": "success"}')
        self.assertEqual(record["device"], DEVICE_ID)

    def test_sensor_log_export_works(self):
        SensorLog.objects.create(device=self.device, status="ONLINE", endpoint_url="x", response={"code": "0"})
        rows = self.export_csv("/admin/core/sensorlog/export/")
        self.assertEqual(rows[0], ["id", "time", "status", "endpoint_url", "device", "response"])
        self.assertEqual(rows[1][2:], ["ONLINE", "x", DEVICE_ID, '{"code": "0"}'])

    def create_recap_data(self):
        for event_id, when in [("a", "2026-09-29T10:00:00+07:00"), ("b", "2026-09-30T10:00:00+07:00"),
                               ("c", "2026-10-01T10:00:00+07:00"), ("d", "2026-10-01T23:30:00+07:00")]:
            EventLog.objects.create(id=event_id, time=when, device=self.device, counted=True)

    def test_recap_date_filter(self):
        self.create_recap_data()
        response = self.client.get("/admin/core/dailyrecap/?date_from=2026-09-30&date_to=2026-10-01")
        days = [row["day"].isoformat() for row in response.context["rows"]]
        self.assertEqual(days, ["2026-10-01", "2026-09-30"])
        self.assertContains(response, "date_from=2026-09-30&amp;date_to=2026-10-01&amp;export=csv")

    def test_recap_export_csv(self):
        self.create_recap_data()
        response = self.client.get("/admin/core/dailyrecap/?date_from=2026-09-30&export=csv")
        self.assertIn("attachment;", response["Content-Disposition"])
        rows = list(csv.reader(io.StringIO(response.content.decode())))
        self.assertEqual(rows[0], ["Date", "Events received", "Counted (in)", "Work Orders sent", "Success", "Failed"])
        self.assertEqual(rows[1:], [["2026-10-01", "2", "2", "0", "0", "0"], ["2026-09-30", "1", "1", "0", "0", "0"]])

    def test_recap_export_xlsx_includes_every_day_without_filter(self):
        from openpyxl import load_workbook

        from core import admin as core_admin

        self.create_recap_data()
        with mock.patch.object(core_admin, "RECAP_DAYS", 1):
            page_rows = self.client.get("/admin/core/dailyrecap/").context["rows"]
            response = self.client.get("/admin/core/dailyrecap/?export=xlsx")
        self.assertEqual(len(page_rows), 1)
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual([row[0] for row in sheet.iter_rows(min_row=2, values_only=True)],
                         ["2026-10-01", "2026-09-30", "2026-09-29"])


@override_settings(ZK_CLIENT_ID="cid", ZK_CLIENT_SECRET="secret")
class BackfillTests(TestCase):
    def test_backfill_stores_history_without_touching_count(self):
        from datetime import datetime

        from django.utils import timezone

        device = DeviceList.objects.get(id=DEVICE_ID)
        device.current_count = 5
        device.save()
        EventLog.objects.create(id="known", time="2026-10-01T11:00:00+07:00", device=device,
                                event_type="out", recognition_target="Cross Line", counted=True)

        def ev(event_id, hhmm, event_type):
            item = event(event_id, event_type)
            item["eventTime"] = f"{hhmm}:00"
            return item

        items = [
            ev("known", "2026-10-01 11:00", "out"),
            ev("a", "2026-10-01 08:00", "in"),
            ev("b", "2026-10-01 07:00", "out"),
            ev("old", "2026-09-30 23:59", "in"),
        ]
        start = timezone.make_aware(datetime(2026, 10, 1))
        with mock.patch.object(services.requests, "get", side_effect=[page(items)]), \
                mock.patch.object(services.requests, "post", side_effect=[token_ok()]) as post:
            stored, relabeled = services.backfill_events(services.ZKClient(), device, start)

        device.refresh_from_db()
        self.assertEqual((stored, relabeled), (2, 1))
        self.assertEqual(device.current_count, 5)
        self.assertEqual(post.call_count, 1)  # token only, no Work Order
        self.assertFalse(EventLog.objects.filter(id="old").exists())
        self.assertEqual(
            sorted(EventLog.objects.filter(counted=True).values_list("id", flat=True)), ["a"]
        )


class EventLogTimeRangeFilterTests(TestCase):
    def test_filters_events_by_time_range(self):
        from django.contrib.auth.models import User

        device = DeviceList.objects.get(id=DEVICE_ID)
        for event_id, hour in (("early", 8), ("mid", 10), ("late", 12)):
            EventLog.objects.create(id=event_id, time=f"2026-10-01T{hour:02d}:00:00+07:00", device=device)
        self.client.force_login(User.objects.create_superuser("admin", "", "pw"))

        response = self.client.get(
            "/admin/core/eventlog/",
            {"time_from": "2026-10-01T09:00", "time_to": "2026-10-01T11:00", "event_type": ""},
        )
        self.assertEqual(response.status_code, 200)
        ids = [obj.id for obj in response.context["cl"].result_list]
        self.assertEqual(ids, ["mid"])
